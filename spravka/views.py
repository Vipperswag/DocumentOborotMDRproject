from django.shortcuts import render, redirect
from django.contrib.auth.forms import UserCreationForm
from django.contrib.auth.models import User
from .forms import DocumentUploadForm
from django.conf import settings
from django.http import FileResponse, HttpResponse
import os
from datetime import datetime
import openpyxl
from docx import Document
import io

# Существующие функции регистрации
def glavnaya(request):
    return render(request, 'index.html')

def signupuser(request):
    if request.method == 'GET':
        return render(request, 'signupuser.html', {'form': UserCreationForm()})
    else:
        if request.POST['password1'] == request.POST['password2']:
            form = UserCreationForm(request.POST)
            if form.is_valid():
                user = form.save()
                return redirect('glavnaya')
            else:
                return render(request, 'signupuser.html', {'form': form})
        else:
            return render(request, 'signupuser.html', {'form': UserCreationForm(), 'error': 'Пароли не совпадают'})

def document_processing(request):
    if request.method == 'POST':
        form = DocumentUploadForm(request.POST, request.FILES)
        if form.is_valid():
            try:
                excel_file = request.FILES['excel_file']

                # Проверяем расширение файла
                if not excel_file.name.endswith(('.xlsx', '.xls')):
                    raise ValueError("Неверный формат файла. Загрузите файл Excel (.xlsx или .xls)")

                # Генерируем Word
                doc_stream = generate_word_document(excel_file)

                response = HttpResponse(
                    doc_stream.getvalue(),
                    content_type='application/vnd.openxmlformats-officedocument.wordprocessingml.document'
                )
                response['Content-Disposition'] = 'attachment; filename="generated_spravka.docx"'
                return response

            except Exception as e:
                return render(request, 'document_processing.html', {
                    'form': form,
                    'error': str(e)
                })
    else:
        form = DocumentUploadForm()

    return render(request, 'document_processing.html', {'form': form})

def generate_word_document(excel_file):
    try:
        # Читаем Excel
        wb = openpyxl.load_workbook(filename=io.BytesIO(excel_file.read()))
        sheet = wb.active

        # Получаем данные (пример для вашей структуры)
        data = {
            'Umya': sheet['A2'].value,
            'Kurs': sheet['B2'].value,
            'Spec': sheet['C2'].value,
            'Formeduc': sheet['D2'].value,
            'NachObuch': sheet['E2'].value,
            'KonObuch': sheet['F2'].value,
            'Prikaz': sheet['G2'].value,
            'Depo': sheet['H2'].value,
            'DataPrikza': sheet['I2'].value,
            'NachPrakt': "01.06.2023",  # Можно добавить в Excel
            'KonPrakt': "30.06.2023",   # Можно добавить в Excel
            'DataVid': datetime.now().strftime("%d.%m.%Y"),
            'NumVid': "12345"
        }

        # Проверяем обязательные поля
        required_fields = ['Umya', 'Kurs', 'Spec', 'Formeduc']
        for field in required_fields:
            if not data.get(field):
                raise ValueError(f"Не заполнено обязательное поле: {field}")

        # Загружаем шаблон Word
        template_path = os.path.join(settings.BASE_DIR, 'spravka', 'static', 'templates', 'template.docx')
        if not os.path.exists(template_path):
            raise FileNotFoundError("Шаблон Word не найден")

        doc = Document(template_path)

        # Заменяем плейсхолдеры
        for paragraph in doc.paragraphs:
            for key, value in data.items():
                if value and f'{{{{{key}}}}}' in paragraph.text:
                    paragraph.text = paragraph.text.replace(f'{{{{{key}}}}}', str(value))

        # Сохраняем в поток
        doc_stream = io.BytesIO()
        doc.save(doc_stream)
        doc_stream.seek(0)
        return doc_stream

    except Exception as e:
        raise Exception(f"Ошибка генерации документа: {str(e)}")

def download_excel_template(request):
    try:
        # Путь к шаблону Excel
        template_path = os.path.join(settings.BASE_DIR, 'spravka', 'static', 'templates', 'form_template.xlsx')

        if not os.path.exists(template_path):
            raise FileNotFoundError("Шаблон Excel не найден")

        with open(template_path, 'rb') as f:
            response = HttpResponse(
                f.read(),
                content_type='application/vnd.openxmlformats-officedocument.spreadsheetml.sheet'
            )
            response['Content-Disposition'] = 'attachment; filename="spravka_template.xlsx"'
            return response

    except Exception as e:
        raise Http404(f"Не удалось загрузить шаблон: {str(e)}")